# -*- coding: utf-8 -*-
"""겟머슬 11세대 RPM LED 자이로볼 — 디자인 전달용 [최종 카피 확정본] -> Word(.docx)

기획안과의 차이: 대안/선택지/디자인 방향 설명을 제거하고, 디자이너가 그대로 박는
'확정 카피'만 잠가서 전달한다.
- 기술명: 'RPM 다이내믹 LED 시스템 (RPM Dynamic LED System)'로 통일 (안전 표기)
- 썸네일: 1안 단독 확정
- 가격/할인: 미확정 → '[가격/할인 영역 — 추후 확정]' placeholder
"""
import os
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document()
style = doc.styles['Normal']
style.font.name = 'Malgun Gothic'
style.font.size = Pt(10.5)
style.element.rPr.rFonts.set(qn('w:eastAsia'), 'Malgun Gothic')

BLUE = RGBColor(0x1E, 0x5E, 0xE0)
PURPLE = RGBColor(0x7A, 0x2B, 0xE0)
RED = RGBColor(0xC0, 0x1A, 0x1A)
GREEN = RGBColor(0x14, 0x7A, 0x3A)
GOLD = RGBColor(0xB8, 0x86, 0x0B)
GREY = RGBColor(0x66, 0x66, 0x66)
DARK = RGBColor(0x11, 0x11, 0x11)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


def set_kfont(run, name='Malgun Gothic'):
    run.font.name = name
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts'); rPr.append(rFonts)
    rFonts.set(qn('w:eastAsia'), name)


def para(text='', size=10.5, bold=False, color=None, align=None, space_after=6, italic=False):
    p = doc.add_paragraph()
    if align:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    if text:
        r = p.add_run(text)
        r.bold = bold; r.italic = italic; r.font.size = Pt(size)
        if color:
            r.font.color.rgb = color
        set_kfont(r)
    return p


def rich(parts, size=10.5, align=None, space_after=6):
    p = doc.add_paragraph()
    if align:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    for text, opt in parts:
        r = p.add_run(text)
        r.bold = opt.get('bold', False); r.italic = opt.get('italic', False)
        r.font.size = Pt(opt.get('size', size))
        if opt.get('color'):
            r.font.color.rgb = opt['color']
        set_kfont(r)
    return p


def bullet(text, color=None, bold=False, size=10.5):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text); r.bold = bold; r.font.size = Pt(size)
    if color:
        r.font.color.rgb = color
    set_kfont(r)
    return p


def heading(text, level=1, color=None):
    h = doc.add_heading(level=level)
    r = h.add_run(text); set_kfont(r)
    r.font.color.rgb = color if color else DARK
    return h


def hr():
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement('w:pBdr'); bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single'); bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:space'), '1'); bottom.set(qn('w:color'), 'CCCCCC')
    pbdr.append(bottom); pPr.append(pbdr)


def box(lines, fill='F4F4F6', border='DDDDDD', size=10.5, color=DARK, bold=True, space_after=8):
    pbox = doc.add_paragraph()
    pbox.paragraph_format.space_after = Pt(space_after)
    pbox.paragraph_format.space_before = Pt(3)
    pPr = pbox._p.get_or_add_pPr()
    shd = OxmlElement('w:shd'); shd.set(qn('w:val'), 'clear'); shd.set(qn('w:fill'), fill); pPr.append(shd)
    bdr = OxmlElement('w:pBdr')
    for side in ('top', 'bottom', 'left', 'right'):
        e = OxmlElement('w:' + side)
        e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), '6'); e.set(qn('w:space'), '6'); e.set(qn('w:color'), border)
        bdr.append(e)
    pPr.append(bdr)
    for i, line in enumerate(lines):
        r = pbox.add_run(line); r.font.size = Pt(size); r.bold = bold
        if color:
            r.font.color.rgb = color
        set_kfont(r)
        if i != len(lines) - 1:
            r.add_break()
    return pbox


def label(text):
    para(text, bold=True, color=GREY, size=8.5, space_after=1)


def set_cell(cell, text, bold=False, color=None, size=9.5, fill=None, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.paragraphs[0].alignment = align
    cell.paragraphs[0].paragraph_format.space_after = Pt(2)
    cell.paragraphs[0].paragraph_format.space_before = Pt(2)
    r = cell.paragraphs[0].add_run(text); r.bold = bold; r.font.size = Pt(size)
    if color:
        r.font.color.rgb = color
    set_kfont(r)
    if fill:
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement('w:shd'); shd.set(qn('w:val'), 'clear'); shd.set(qn('w:fill'), fill); tcPr.append(shd)


def make_table(rows, highlight_col=None):
    ncols = len(rows[0])
    table = doc.add_table(rows=len(rows), cols=ncols)
    table.style = 'Table Grid'; table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = table.cell(ri, ci)
            if ri == 0:
                set_cell(cell, val, bold=True, color=WHITE, fill='1F2937', align=WD_ALIGN_PARAGRAPH.CENTER)
            else:
                hl = (highlight_col is not None and ci == highlight_col)
                set_cell(cell, val, bold=hl, color=(GOLD if hl else None), fill=('FFF6E0' if hl else None))
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


TECH = 'RPM 다이내믹 LED 시스템 (RPM Dynamic LED System)'
PRICE = '[가격/할인 영역 — 추후 확정]'

# ===== 표지 =====
t = doc.add_heading(level=0); tr = t.add_run('겟머슬 11세대 RPM LED 자이로볼'); set_kfont(tr)
s = doc.add_heading(level=1); sr = s.add_run('디자인 전달용 — 최종 카피 확정본'); set_kfont(sr); sr.font.color.rgb = DARK

box([
    '✔ 이 문서의 문구는 모두 확정본입니다. 디자이너는 시각 구성에만 집중하세요.',
    '✔ 표기된 텍스트는 그대로 사용하며, 임의 수정·카피라이팅이 필요 없습니다.',
    '✔ 문구 변경이 필요하면 디자인 진행 전 기획에 회신 부탁드립니다.',
], fill='EAF2FF', border='9CC0FF', size=10, color=DARK, bold=False)

label('확정 기준 (이번 버전에서 잠근 항목)')
bullet('기술명 표기: "RPM 다이내믹 LED 시스템 (RPM Dynamic LED System)"으로 통일', color=GREEN)
bullet('썸네일: 1안(LED 변화 강조) 단독 확정', color=GREEN)
bullet('가격·할인·쿠폰 수치: 미확정 → 배지/CTA에 "[가격/할인 영역 — 추후 확정]" 공간만 확보', color=GREEN)
para('※ 실제 별도 센서칩이 제품에 탑재돼 있다면 기술명을 "RPM 다이내믹 센서 칩"으로 되돌려도 됩니다(기획 확인 필요). '
     '금지/대체 표현은 마지막 페이지 참고.', size=9, color=RED, italic=True)
hr()

# ===== 15장 확정 카피 =====
heading('상세페이지 15장 — 확정 카피', 1)

# 1
heading('1장. 메인 후킹', 2, color=BLUE)
label('헤드라인'); box(['손목은 매일 쓰는데, 손목 운동은 왜 안 하세요?'], size=12)
label('서브카피'); box(['마우스, 스마트폰, 헬스, 골프, 테니스까지', '하루 종일 버티는 손목을 위한 11세대 자이로 손목 운동기구'], size=10.5, bold=False)
label('포인트 배지 (5개)')
for b in ['회전력 따라 LED 변화', '손목·악력·전완근 자극', '태엽식 스타트', '스트랩 + 전용케이스 증정', '30일 무료체험']:
    bullet(b, color=GREEN, bold=True)
label('가격 영역'); para(PRICE, color=RED, bold=True)

# 2
heading('2장. 문제 공감', 2, color=BLUE)
label('헤드라인'); box(['당신의 손목, 생각보다 더 많이 일하고 있습니다'], size=12)
label('본문 (4줄 + 마무리)')
for b in ['헬스할 때 손목이 먼저 흔들린다면', '골프·테니스 후 손목이 쉽게 피로하다면',
          '마우스와 키보드를 오래 써 손이 뻐근하다면', '일반 악력기가 지루해서 오래 못 했다면']:
    bullet(b)
box(['이제 손목도 따로 운동할 때입니다.'], size=11)

# 3
heading('3장. 제품 정의', 2, color=BLUE)
label('헤드라인'); box(['이건 단순히 반짝이는 공이 아닙니다'], size=12)
label('본문'); para('겟머슬 자이로볼은 손 안에서 회전 저항을 만들고, 그 저항을 손목으로 버티며 운동하는 자이로 손목 트레이너입니다.')
label('강조 한 줄'); box(['작다. 간단하다. 그런데 손목은 바로 압니다.'], size=11, color=PURPLE)

# 4
heading('4장. 구형 제품 비교', 2, color=BLUE)
label('헤드라인'); box(['구형 1~10세대와 비교하지 마세요'], size=12)
make_table([
    ['구형 자이로볼', '겟머슬 11세대'],
    ['랜덤으로만 번쩍이는 LED', '회전력 따라 LED 변화'],
    ['의미 없는 단순 점멸', 'RPM 다이내믹 LED 시스템'],
    ['그냥 돌아가는 장난감 느낌', '손목으로 버티는 자이로 저항감'],
    ['금방 질리는 사용감', 'LED 변화로 운동 몰입감 상승'],
    ['초반 회전 잡기가 번거로움', '태엽식 스타트'],
    ['본품만 있는 구성', '스트랩 + 전용케이스 증정'],
    ['써보기 전 운동감이 불안', '30일 무료체험 보장'],
], highlight_col=1)
label('하단 카피'); box(['구형은 그냥 번쩍였습니다 — 겟머슬은 회전력에 따라 반응합니다'], size=10.5, color=PURPLE)

# 5
heading('5장. RPM 다이내믹 LED 시스템', 2, color=BLUE)
label('헤드라인'); box(['랜덤 LED가 아닙니다'], size=12)
label('기술명 (확정 표기)'); box([TECH], size=11, color=GOLD)
label('본문'); para('회전력에 따라 LED 반응이 달라집니다. 느리게 돌릴 때는 은은하게, 속도가 올라갈수록 선명하게, 강하게 회전할수록 더 역동적으로.')
label('강조 한 줄'); box(['내 손목 스피드가 불빛으로 보입니다'], size=11, color=PURPLE)

# 6
heading('6장. 태엽식 스타트', 2, color=BLUE)
label('헤드라인'); box(['태엽 감고 바로 시작하세요'], size=12)
label('본문'); para('줄을 찾고 감고 실패하는 번거로움 대신, 태엽을 감아 초반 회전을 시작합니다. 그다음은 내 손목으로 속도를 키우고 자이로 저항을 버티는 방식입니다.')
label('사용 순서 (5단계 — 아이콘+숫자)')
make_table([['1', '2', '3', '4', '5'],
            ['태엽을 감는다', '회전을 시작한다', '손목 스냅으로 속도 ↑', 'LED가 살아난다', '저항감이 강해진다']])
label('하단 카피'); box(['시작은 쉽게, 운동감은 직접적으로'], size=10.5, color=PURPLE)

# 7
heading('7장. 진짜 운동감', 2, color=BLUE)
label('헤드라인'); box(['자동으로 도는 장난감이 아닙니다'], size=12)
label('본문'); para('겟머슬은 스스로 계속 돌아가는 제품이 아닙니다. 내 손목 움직임으로 회전을 유지하고, 내 힘으로 저항을 키우는 진짜 자이로 운동기구입니다.')
label('강조 한 줄'); box(['내가 돌린 만큼 저항이 커지고, 내가 컨트롤한 만큼 손목이 반응합니다.'], size=11, color=PURPLE)

# 8
heading('8장. 악력기·덤벨 비교', 2, color=BLUE)
label('헤드라인'); box(['악력기는 쥐기만 합니다 — 자이로볼은 손목이 계속 반응합니다'], size=11.5)
make_table([
    ['구분', '악력기', '덤벨', '겟머슬 자이로볼'],
    ['운동 방식', '쥐기 반복', '무게 들기', '회전 저항 유지'],
    ['재미', '단조로움', '보통', 'LED + 회전감'],
    ['휴대성', '좋음', '낮음', '좋음'],
    ['손목 컨트롤', '낮음', '보통', '높음'],
    ['전완근 자극', '가능', '가능', '회전 유지형 자극'],
    ['몰입감', '낮음', '보통', '높음'],
], highlight_col=3)
label('하단 카피'); box(['쥐는 힘만이 아니라, 손목으로 버티는 힘까지'], size=10.5, color=PURPLE)

# 9
heading('9장. 운동 부위', 2, color=BLUE)
label('헤드라인'); box(['손목부터 전완근까지, 손 안에서 자극'], size=12)
label('본문'); para('회전을 유지할수록 자이로 저항감이 강해집니다. 그 저항을 손목으로 버티며 손목·악력·전완근 운동에 도움을 줄 수 있습니다.')
label('강조 부위 (도해 라벨)'); para('손목 · 악력 · 전완근 · 팔 회전 감각', bold=True, color=GREEN)

# 10
heading('10장. 타깃별 사용 장면', 2, color=BLUE)
label('헤드라인'); box(['손목을 많이 쓰는 사람이라면 하나쯤 필요합니다'], size=12)
make_table([
    ['타깃', '제안 루틴'],
    ['헬스', '덤벨·바벨·턱걸이 전 손목과 전완근을 깨우는 루틴'],
    ['골프', '스윙 전 손목 감각과 전완근 컨트롤 루틴'],
    ['테니스·배드민턴', '라켓을 잡는 손목 힘과 회전 감각 관리'],
    ['클라이밍', '손가락·손목·전완근을 함께 쓰는 감각 훈련'],
    ['직장인·게이머', '마우스·키보드로 지친 손목을 위한 짧은 루틴'],
    ['선물용', '운동 좋아하는 사람에게 부담 없는 실용 선물'],
])

# 11
heading('11장. 하루 10분 루틴', 2, color=BLUE)
label('헤드라인'); box(['처음이라면 하루 1분부터 시작하세요'], size=12)
make_table([
    ['단계', '시간', '방법', '목적'],
    ['Lv.1 적응', '30초~1분', '천천히 회전 유지', '감각 익히기'],
    ['Lv.2 유지', '2~3분', '일정 속도 유지', '손목 컨트롤'],
    ['Lv.3 강화', '3~5분', '속도 올리기', '전완근 자극'],
    ['Lv.4 루틴', '5~10분', '양손 번갈아 사용', '홈트 루틴'],
])
label('하단 카피'); para('처음부터 오래 하지 마세요. 짧게 시작하고, 익숙해질수록 속도와 시간을 늘려보세요.', color=GREY)

# 12
heading('12장. 단일 구성 + 사은품', 2, color=BLUE)
label('헤드라인'); box(['옵션 고민 없이, 필요한 구성만 한 번에'], size=12)
label('구성품 (3종)')
for b in ['겟머슬 11세대 RPM LED 자이로볼 본품', '안전 스트랩 증정', '전용 보관 케이스 증정']:
    bullet(b, color=GREEN, bold=True)
label('본문'); para('고속 회전 시 안정감을 더해주는 스트랩, 보관과 휴대를 편하게 해주는 전용 케이스까지. 사자마자 바로 쓰고, 쓰고 나면 깔끔하게 보관하세요.')
label('배너'); box(['본품 + 스트랩 + 전용케이스 — 단일 구성으로 바로 시작'], size=10.5)
para('※ 색상 옵션·컬러 선택 문구는 넣지 않습니다 (단일 제품 구성).', size=9, color=RED, italic=True)

# 13
heading('13장. 광고비 대신 제품력', 2, color=BLUE)
label('헤드라인'); box(['비싼 자이로볼 가격, 정말 제품값일까요 광고값일까요?'], size=11.5)
label('본문'); para('겟머슬은 유명인 광고비를 제품 가격에 얹지 않았습니다. 그 비용 대신 11세대 구조, 태엽식 스타트, RPM 다이내믹 LED 시스템, 자이로 저항감, 스트랩·전용 케이스 구성에 집중했습니다.')
label('강조 카피'); box(['광고비는 빼고, 제품력은 채운 11세대 자이로볼'], size=11, color=PURPLE)

# 14
heading('14장. 30일 무료체험 보장', 2, color=BLUE)
label('헤드라인'); box(['뜯어서 30일 동안 마음껏 돌려보세요'], size=12)
label('서브'); para('제품에 자신 있으니까 약속합니다', bold=True)
label('체크 포인트 (3개)')
for b in ['개봉해도 OK', '사용해도 OK', '단순변심도 OK']:
    bullet(b, color=GREEN, bold=True)
label('최종 후킹 (보증 배지)'); box(['만족하지 못하면 100% 무료 환불'], size=12, color=GOLD)
label('본문'); para('사진만 보고는 자이로볼의 운동감을 알 수 없습니다. 직접 돌려보고, LED 반응을 보고, 손목에 전해지는 저항감을 느껴보세요. 만족하지 못하셨다면 30일 이내 100% 무료 환불해드립니다.')
label('하단 안내 문구 (작게)')
para('※ 30일 무료체험 및 무료 환불은 상품 수령일 기준 30일 이내 신청 건에 한해 적용됩니다. 구성품 누락, 고의 훼손, 사용 불가 수준의 파손이 있는 경우 제한될 수 있습니다.',
     size=8.5, color=GREY)

# 15
heading('15장. FAQ + 최종 클로징', 2, color=BLUE)
label('헤드라인'); box(['구매 전 궁금한 점, 미리 확인하세요'], size=12)
faqs = [
    ('Q. 소리가 나나요?', '네. 내부 회전체가 고속으로 회전하는 운동기구 특성상 회전 소리와 진동이 발생할 수 있으며, 이는 정상적인 작동 특성입니다.'),
    ('Q. 자동으로 계속 돌아가나요?', '아니요. 태엽으로 초반 회전을 시작한 뒤, 사용자의 손목 움직임으로 회전을 유지하는 제품입니다.'),
    ('Q. 초보자도 사용할 수 있나요?', '처음에는 30초~1분씩 천천히 시작하는 것을 권장합니다. 익숙해질수록 회전 유지 감각이 좋아집니다.'),
    ('Q. LED는 어떻게 나오나요?', '회전력과 속도에 따라 LED 반응이 달라지는 구조입니다. 손목 스피드에 따라 밝기·연출감이 다르게 나타날 수 있습니다.'),
    ('Q. 손목 통증 치료용인가요?', '아니요. 본 제품은 의료기기가 아닌 운동 보조용품입니다. 통증·질환이 있는 경우 전문가 상담 후 사용해 주세요.'),
]
for q, a in faqs:
    para(q, bold=True, color=BLUE, space_after=1); para(a, space_after=6)
label('최종 클로징 (다크 박스)')
box(['저가형 자이로볼에 실망했다면, 이제 11세대로 바꾸세요', '',
     '랜덤 LED가 아닙니다. 회전력에 따라 LED가 달라집니다.',
     '자동 장난감이 아닙니다. 내 손목으로 저항을 키우는 운동기구입니다.',
     '광고비로 가격을 부풀리지 않았습니다. 제품력, 사은품, 30일 무료체험으로 증명합니다.'],
    fill='111418', border='111418', size=10.5, color=WHITE)
label('최종 CTA'); box(['광고비는 빼고, 제품력은 채우고, 30일 무료체험까지 — 겟머슬로 손목 운동 루틴을 시작하세요'], size=10.5, color=GOLD)
label('CTA 가격 영역'); para(PRICE, color=RED, bold=True)
hr()

# ===== 썸네일 확정 =====
heading('썸네일 — 확정안 (1안 단독)', 1)
box(['회전력 따라 LED 변화', '11세대 RPM LED 자이로볼', '손목·악력·전완근 자극'],
    fill='FFF6E0', border='E5C76B', size=11, color=DARK)
para('※ 2~5안은 이번 디자인에서 사용하지 않습니다(메인 썸네일은 1안 단독 확정).', size=9, color=GREY, italic=True)
hr()

# ===== 반복 핵심 문장 =====
heading('페이지 전반 반복 문장 (확정)', 1)
para('아래 문장은 페이지 곳곳에 반복 배치합니다.', color=GREY)
for b in ['랜덤 LED 구형은 그만', '회전력 따라 LED가 달라지는 11세대 자이로볼',
          '자동으로 도는 장난감이 아닙니다', '내 손목으로 저항을 키우는 운동기구입니다',
          '속도감은 눈으로, 저항감은 손목으로', '악력기는 쥐기만 합니다. 자이로볼은 손목이 계속 반응합니다',
          '광고비는 빼고, 제품력은 채웠습니다', '뜯어서 30일 동안 마음껏 돌려보세요',
          '개봉해도 OK, 사용해도 OK, 단순변심도 OK', '만족하지 못하면 100% 무료 환불']:
    bullet(b, bold=True)
hr()

# ===== 금지 / 대체 표현 (레퍼런스) =====
heading('금지 표현 / 대체 표현 (반드시 준수)', 1)
make_table([
    ['❌ 절대 사용 금지', '✅ 대신 사용'],
    ['AI 칩 / 인공지능 LED / AI 회전 분석', 'RPM 다이내믹 LED 시스템 / 회전력 반응형 LED'],
    ['무소음 / 소음 없음 / 완전 무진동', '자이로볼 특유의 회전 소리와 진동(정상 특성)'],
    ['자동 회전 보정 / 자동 궤도 보정', '내 손목으로 직접 유지하는 회전 운동'],
    ['고장 절대 없음 / 무조건 100% 가능', '돌릴수록 강해지는 자이로 저항감'],
    ['손목 통증 치료 / 재활 / 관절 개선 / 염증 완화', '손목·악력·전완근 자극 / 손목 컨트롤 훈련'],
    ['근육 증가 보장 / 악력 무조건 증가', '운동에 도움을 줄 수 있습니다 (단정 X)'],
    ['의료기기처럼 보이는 표현', '운동 보조용품 / 10분 손목 루틴'],
])
para('※ "센서 칩" 단독 표기는 실제 칩 부품 미탑재 시 허위표시 리스크가 있어 "RPM 다이내믹 LED 시스템"으로 통일했습니다.',
     size=9, color=RED, italic=True)

out = os.path.join(os.path.dirname(__file__), '..', 'docs', '겟머슬_최종카피_확정본.docx')
out = os.path.abspath(out)
doc.save(out)
print('saved:', out)
