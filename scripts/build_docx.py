# -*- coding: utf-8 -*-
"""겟머슬 인플루언서 협업 브리프 -> Word(.docx) 생성"""
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document()

# 기본 폰트(한글) 설정
style = doc.styles['Normal']
style.font.name = 'Malgun Gothic'
style.font.size = Pt(10.5)
style.element.rPr.rFonts.set(qn('w:eastAsia'), 'Malgun Gothic')

BLUE = RGBColor(0x1E, 0x5E, 0xE0)
PURPLE = RGBColor(0x7A, 0x2B, 0xE0)
RED = RGBColor(0xC0, 0x1A, 0x1A)
GREEN = RGBColor(0x14, 0x7A, 0x3A)
GREY = RGBColor(0x66, 0x66, 0x66)


def set_kfont(run, name='Malgun Gothic'):
    run.font.name = name
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:eastAsia'), name)


def para(text='', size=10.5, bold=False, color=None, align=None, space_after=6, italic=False):
    p = doc.add_paragraph()
    if align:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    if text:
        r = p.add_run(text)
        r.bold = bold
        r.italic = italic
        r.font.size = Pt(size)
        if color:
            r.font.color.rgb = color
        set_kfont(r)
    return p


def rich(parts, size=10.5, align=None, space_after=6):
    """parts: list of (text, {bold,color,italic})"""
    p = doc.add_paragraph()
    if align:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    for text, opt in parts:
        r = p.add_run(text)
        r.bold = opt.get('bold', False)
        r.italic = opt.get('italic', False)
        r.font.size = Pt(opt.get('size', size))
        if opt.get('color'):
            r.font.color.rgb = opt['color']
        set_kfont(r)
    return p


def bullet(text, color=None, bold=False):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text)
    r.bold = bold
    r.font.size = Pt(10.5)
    if color:
        r.font.color.rgb = color
    set_kfont(r)
    return p


def heading(text, level=1):
    h = doc.add_heading(level=level)
    r = h.add_run(text)
    set_kfont(r)
    if level == 1:
        r.font.color.rgb = RGBColor(0x11, 0x11, 0x11)
    return h


def hr():
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), 'CCCCCC')
    pbdr.append(bottom)
    pPr.append(pbdr)


# ===== 표지 / 타이틀 =====
title = doc.add_heading(level=0)
tr = title.add_run('겟머슬 11세대 RPM LED 자이로볼')
set_kfont(tr)
sub = doc.add_heading(level=1)
sr = sub.add_run('운동 인플루언서 협업 브리프')
set_kfont(sr)
para('※ 콘티를 그대로 따라 찍는 광고가 아닙니다. 평소 운동 루틴 콘텐츠 스타일 그대로, '
     '제품을 자연스럽게 녹여 소개해 주세요. 아래 [필수 표기]·[금지 표현]만 지켜주시면 '
     '연출·편집·말투·길이는 자유입니다.', size=9.5, color=GREY, italic=True)
hr()

# 0. 한 줄 요약
heading('0. 협업 한 줄 요약', 2)
para('손목·악력·전완근을 쓰는 운동인에게, "자동으로 도는 장난감이 아니라 내 손목으로 저항을 '
     '키우는 11세대 자이로 손목 운동기구"를 본인의 운동 루틴 안에서 자연스럽게 소개.')

# 1. 제품 이해
heading('1. 이 제품이 뭔가요 (핵심 이해)', 2)
bullet('겟머슬 11세대 RPM LED 자이로볼 — 손목/악력/전완근 자극용 자이로볼')
bullet('회전력에 따라 LED가 달라짐 (느릴 땐 은은, 빠를수록 선명·역동) — 구형 랜덤 점멸과 다름')
bullet('태엽 감아 시작 → 그다음은 내 손목 스냅으로 회전 유지, 돌릴수록 저항감↑')
bullet('자동으로 계속 도는 제품 아님. 내 힘으로 하는 운동기구')
bullet('구성: 본품 + 안전 스트랩 + 전용 케이스 / 30일 무료체험 · 100% 무료 환불')

# 2. 꼭 살릴 3가지
heading('2. 영상에서 꼭 살려주실 3가지 (본인 말투로)', 2)
rich([('① ', {'bold': True, 'color': BLUE}),
      ('"이거 자동으로 도는 장난감 아니에요. 내 손목으로 돌리고 버티는 거예요."', {'bold': True})])
rich([('② ', {'bold': True, 'color': BLUE}),
      ('회전력에 따라 LED가 달라진다 — 직접 천천히/빠르게 돌려서 차이를 보여주기', {'bold': True})])
rich([('③ ', {'bold': True, 'color': BLUE}),
      ('구성·혜택 — 스트랩+케이스 증정, 30일 무료체험, 만족 못하면 100% 환불', {'bold': True})])
para('대본 없이 본인 경험·느낌으로 말해주시는 게 제일 좋습니다. '
     '(예: "세트 사이에 손목 풀 때 쓰는데 전완근 펌핑 장난 아님")', size=9.5, color=GREY, italic=True)

# 3. 자연스러운 노출 아이디어
heading('3. 자연스러운 노출 아이디어 (택1~2, 자유)', 2)
for t in [
    '운동 전 워밍업 루틴에 끼워넣기: "본 운동 전 손목·전완근 깨우는 루틴 하나 추가했어요"',
    '세트 사이 인터벌: 휴식 때 손목 자극 → "쉬는 시간에 손목 안 놀려요"',
    '데스크/일상 컷: 키보드·마우스로 손목 뻐근할 때 잠깐 돌리기 (직장인·게이머 공감)',
    '종목 연계: 골프/테니스/클라이밍 등 본인 종목의 "손목 힘"과 연결',
    '솔직 리뷰 톤: "저가 자이로볼 써봤는데 금방 질렸거든요? 이건 LED가 회전 따라 변해서…"',
    '만족감 컷: 회전 소리 + LED 잔상 슬로우모션 한 컷 (있으면 도달에 좋음, 필수 아님)',
]:
    bullet(t)

# 4. 권장 비주얼
heading('4. 보여주시면 좋은 비주얼 (권장, 강제 아님)', 2)
for t in [
    '손에 쥐고 천천히 → 빠르게 돌리며 LED 변화 보여주기',
    '태엽 감아 시작하는 장면 / 손목 스냅으로 속도 올리는 장면',
    '돌린 뒤 전완근 펌핑/뻐근함 직접 코멘트',
    '마지막에 본품+스트랩+케이스 한 번 보여주기',
]:
    bullet(t)

# 5. 톤/길이/포맷
heading('5. 톤 / 길이 / 포맷', 2)
bullet('본인 평소 스타일·길이 그대로 (릴스/쇼츠 15~45초 권장, 강제 아님)')
bullet('9:16 세로, 평소 쓰시는 화질이면 충분')
bullet('분위기: 솔직·자연스럽게. 과장 광고 톤 지양')

hr()

# 6. 필수 표기
heading('6. ★ 필수 표기 (꼭 넣어주세요)', 2)
rich([('유료 광고/협찬 표기 의무: ', {'bold': True, 'color': RED}),
      ('영상 내 자막 + 캡션에 #광고 또는 #유료광고 (브랜드 협찬임을 명확히)', {})])
bullet('인스타: 상단 "유료 파트너십" 라벨 + 캡션 #광고')
bullet('유튜브: "유료 프로모션 포함" 체크 + 설명란 명시')
bullet('계정 태그 / 구매링크: (브랜드 계정·구매링크 — 추후 전달)')
bullet('엔드/캡션에 한 번은: 겟머슬 11세대 RPM LED 자이로볼 + 30일 무료체험 / 100% 무료 환불')

# 7. 금지 표현
heading('7. ⛔ 금지 표현 (법적 안전 — 반드시 지켜주세요)', 2)
para('효능·의료 표현, 절대 단정 표현은 사용 금지입니다.', bold=True)
for t in [
    '손목 통증 치료 / 재활 / 관절 개선 / 염증 완화 (→ 의료 효능 표현 금지)',
    '근육 증가 보장 / 악력 무조건 증가 / 초보자도 무조건 100%',
    '무소음 / 소음 없음 / 완전 무진동 (자이로볼은 회전 소리·진동이 정상)',
    'AI 칩 / 인공지능 LED / 자동 회전 보정 / 고장 절대 없음',
]:
    bullet(t, color=RED)
para('✅ 대신 이렇게:', bold=True, color=GREEN)
for t in [
    '"손목·악력·전완근 자극", "도움을 줄 수 있어요" (단정 X)',
    '"회전 소리랑 진동이 좀 있는데 그게 운동되는 느낌" (소리는 정상 특성으로 안내)',
    '기술명은 "RPM 다이내믹 센서 칩" / "회전력 반응형 LED" 로 표현',
]:
    bullet(t, color=GREEN)

hr()

# 8. 해시태그
heading('8. 해시태그 예시 (참고)', 2)
para('#겟머슬  #자이로볼  #손목운동  #전완근  #악력  #홈트  #운동기구  #광고')

# 9. 제공/준비
heading('9. 제공 / 준비', 2)
bullet('제품 제공: 본품 + 스트랩 + 전용케이스 (실사용 후 솔직 소개)')
bullet('브랜드 계정·구매링크·해시태그 최종본: 담당자 추후 전달')
bullet('초안 컷/캡션은 업로드 전 가볍게 컨펌(필수 표기·금지 표현만 체크), 창작은 자유')

hr()

# ===== 10. 채널 맞춤 제안 =====
heading('10. 이 채널에 딱 맞는 제안 (맨몸운동 · 홈트 · 헬린이)', 2)
rich([('핵심 연결고리  ', {'bold': True, 'color': PURPLE}),
      ('푸쉬업·플랭크 같은 맨몸운동은 손목·전완근 부하가 큽니다. '
       '"본 운동 전 손목 워밍업 루틴"으로 자이로볼을 끼우면 채널 콘텐츠와 100% 자연스럽게 연결돼요. '
       '(헬린이 타깃에 특히 잘 맞음)', {})])
para('추천 포맷 (본인 시그니처 그대로):', bold=True)
bullet('가슴/푸쉬업 루틴 영상 도입부 또는 마무리에 "손목 워밍업" 한 코너로 자이로볼 추가')
bullet('"맨몸운동 많이 하면 손목에 부하 쌓이죠? 저는 본 운동 전에 이걸로 손목부터 깨워요" 식 솔직 멘트')
bullet('평소처럼 동작 넘버링(1️⃣2️⃣3️⃣) + 📌포인트 + 🔖저장 유도 포맷 유지')
para('⚠️ 주의: "통증/부상 예방·치료"로 말하지 말고 "워밍업·준비운동·자극"으로만 표현해 주세요. '
     '(의료 효능 표현은 금지 — 7번 참고)', size=9.5, color=RED, italic=True)

# ===== 11. 바로 쓰는 캡션 예시 (이분 스타일) =====
heading('11. 바로 쓰는 캡션 예시 (이분 스타일 그대로)', 2)
para('아래는 참고용 초안입니다. 본인 말투로 자유롭게 바꿔주세요.', size=9.5, color=GREY, italic=True)

caption_lines = [
    '💪 푸쉬업 많이 하는 헬린이 주목',
    '가슴 루틴 전, 손목부터 깨우세요 🔥',
    '',
    '맨몸운동 많이 하면 손목·전완근에 은근 부하가 쌓여요.',
    '저는 본 운동 전에 이걸로 손목 워밍업부터 합니다.',
    '',
    '🌀 겟머슬 11세대 RPM LED 자이로볼',
    '1️⃣ 태엽 감고 천천히 회전 시작 (30초)',
    '2️⃣ 손목 스냅으로 속도 올리기 (회전 따라 LED가 달라져요)',
    '3️⃣ 양손 번갈아 1분씩',
    '',
    '📌 포인트',
    '자동으로 도는 장난감 아님 → 내 손목으로 돌리고 버티는 운동',
    '돌릴수록 자이로 저항감이 강해져서 전완근까지 자극돼요',
    '',
    '🔖 저장해두고 푸쉬업 루틴 전 손목 워밍업으로 써보세요',
    '(스트랩+케이스 증정 / 30일 무료체험 / 만족 못하면 100% 무료 환불)',
    '',
    '#광고 #겟머슬 #자이로볼 #손목운동 #전완근 #맨몸운동 #헬린이 #홈트',
]
# 회색 박스 느낌으로 한 문단에 묶기
pbox = doc.add_paragraph()
pbox.paragraph_format.space_after = Pt(8)
pPr = pbox._p.get_or_add_pPr()
shd = OxmlElement('w:shd'); shd.set(qn('w:val'), 'clear'); shd.set(qn('w:fill'), 'F4F4F6')
pPr.append(shd)
bdr = OxmlElement('w:pBdr')
for side in ('top', 'bottom', 'left', 'right'):
    e = OxmlElement('w:' + side)
    e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), '6'); e.set(qn('w:space'), '6'); e.set(qn('w:color'), 'DDDDDD')
    bdr.append(e)
pPr.append(bdr)
for i, line in enumerate(caption_lines):
    r = pbox.add_run(line)
    r.font.size = Pt(10)
    set_kfont(r)
    if i != len(caption_lines) - 1:
        r.add_break()

hr()
rich([('한 줄 핵심  ', {'bold': True, 'color': PURPLE}),
      ('"내 손목으로 돌리는 운동기구, 회전 따라 LED가 변한다 — 운동 루틴 속에서 솔직하게."',
       {'bold': True, 'italic': True})])

import os
out = os.path.join(os.path.dirname(__file__), '..', 'docs', '겟머슬_인플루언서_협업_브리프.docx')
out = os.path.abspath(out)
doc.save(out)
print('saved:', out)
