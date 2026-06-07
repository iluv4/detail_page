# -*- coding: utf-8 -*-
"""겟머슬 11세대 RPM LED 자이로볼 — 상세페이지 최종 기획안 -> Word(.docx) 생성"""
import os
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document()

# 기본 폰트(한글) 설정
style = doc.styles['Normal']
style.font.name = 'Malgun Gothic'
style.font.size = Pt(10.5)
style.element.rPr.rFonts.set(qn('w:eastAsia'), 'Malgun Gothic')

BLUE = RGBColor(0x1E, 0x5E, 0xE0)
PURPLE = RGBColor(0x7A, 0x2B, 0xE0)
RED = RGBColor(0xC0, 0x1A, 0x1A)
GREEN = RGBColor(0x14, 0x7A, 0x3A)
GOLD = RGBColor(0xB8, 0x86, 0x0B)
GREY = RGBColor(0x66, 0x66, 0x66)
DARK = RGBColor(0x11, 0x11, 0x11)


def set_kfont(run, name='Malgun Gothic'):
    run.font.name = name
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:eastAsia'), name)


def para(text='', size=10.5, bold=False, color=None, align=None, space_after=6, italic=False):
    p = doc.add_paragraph()
    if align:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    if text:
        r = p.add_run(text)
        r.bold = bold
        r.italic = italic
        r.font.size = Pt(size)
        if color:
            r.font.color.rgb = color
        set_kfont(r)
    return p


def rich(parts, size=10.5, align=None, space_after=6):
    """parts: list of (text, {bold,color,italic,size})"""
    p = doc.add_paragraph()
    if align:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    for text, opt in parts:
        r = p.add_run(text)
        r.bold = opt.get('bold', False)
        r.italic = opt.get('italic', False)
        r.font.size = Pt(opt.get('size', size))
        if opt.get('color'):
            r.font.color.rgb = opt['color']
        set_kfont(r)
    return p


def bullet(text, color=None, bold=False, size=10.5):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text)
    r.bold = bold
    r.font.size = Pt(size)
    if color:
        r.font.color.rgb = color
    set_kfont(r)
    return p


def numbered(text, color=None, bold=False):
    p = doc.add_paragraph(style='List Number')
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text)
    r.bold = bold
    r.font.size = Pt(10.5)
    if color:
        r.font.color.rgb = color
    set_kfont(r)
    return p


def heading(text, level=1, color=None):
    h = doc.add_heading(level=level)
    r = h.add_run(text)
    set_kfont(r)
    r.font.color.rgb = color if color else DARK
    return h


def hr():
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), 'CCCCCC')
    pbdr.append(bottom)
    pPr.append(pbdr)


def shade_para(lines, fill='F4F4F6', border='DDDDDD', size=10, color=None, bold=False):
    """회색 박스 안에 여러 줄 묶기 (강조 카피용)"""
    pbox = doc.add_paragraph()
    pbox.paragraph_format.space_after = Pt(8)
    pbox.paragraph_format.space_before = Pt(4)
    pPr = pbox._p.get_or_add_pPr()
    shd = OxmlElement('w:shd'); shd.set(qn('w:val'), 'clear'); shd.set(qn('w:fill'), fill)
    pPr.append(shd)
    bdr = OxmlElement('w:pBdr')
    for side in ('top', 'bottom', 'left', 'right'):
        e = OxmlElement('w:' + side)
        e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), '6'); e.set(qn('w:space'), '6'); e.set(qn('w:color'), border)
        bdr.append(e)
    pPr.append(bdr)
    for i, line in enumerate(lines):
        r = pbox.add_run(line)
        r.font.size = Pt(size)
        r.bold = bold
        if color:
            r.font.color.rgb = color
        set_kfont(r)
        if i != len(lines) - 1:
            r.add_break()
    return pbox


def set_cell(cell, text, bold=False, color=None, size=9.5, fill=None, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.paragraphs[0].alignment = align
    cell.paragraphs[0].paragraph_format.space_after = Pt(2)
    cell.paragraphs[0].paragraph_format.space_before = Pt(2)
    r = cell.paragraphs[0].add_run(text)
    r.bold = bold
    r.font.size = Pt(size)
    if color:
        r.font.color.rgb = color
    set_kfont(r)
    if fill:
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement('w:shd'); shd.set(qn('w:val'), 'clear'); shd.set(qn('w:fill'), fill)
        tcPr.append(shd)


def make_table(rows, header_fill='1F2937', header_color=RGBColor(0xFF, 0xFF, 0xFF),
               highlight_col=None, highlight_fill='FFF6E0'):
    """rows[0] = header. highlight_col: 강조할 열 인덱스(겟머슬 열)."""
    ncols = len(rows[0])
    table = doc.add_table(rows=len(rows), cols=ncols)
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = table.cell(ri, ci)
            if ri == 0:
                set_cell(cell, val, bold=True, color=header_color, size=9.5,
                         fill=header_fill, align=WD_ALIGN_PARAGRAPH.CENTER)
            else:
                is_hl = (highlight_col is not None and ci == highlight_col)
                set_cell(cell, val, bold=is_hl,
                         color=(GOLD if is_hl else None),
                         fill=(highlight_fill if is_hl else None))
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


# =====================================================================
# 표지 / 타이틀
# =====================================================================
title = doc.add_heading(level=0)
tr = title.add_run('👑 겟머슬 11세대 RPM LED 자이로볼')
set_kfont(tr)
sub = doc.add_heading(level=1)
sr = sub.add_run('구매전환 극대화 최종 상세페이지 기획안')
set_kfont(sr)
sr.font.color.rgb = DARK

shade_para([
    '손목은 매일 쓰는데, 손목 운동은 왜 안 하세요?',
    '',
    '마우스, 스마트폰, 헬스, 골프, 테니스, 게임까지 하루 종일 버티는 손목을 위한',
    '겟머슬 11세대 RPM LED 자이로볼',
    '',
    '회전력 따라 LED가 달라지고 · 돌릴수록 자이로 저항감이 강해지며',
    '손목·악력·전완근을 손 안에서 자극합니다',
], fill='111418', border='111418', size=11, color=RGBColor(0xFF, 0xFF, 0xFF), bold=True)

para('벤치마킹 분석상 본 제품은 "강한 남성 운동기구" 느낌은 살리되, 구매층을 넓히려면 '
     '손목을 많이 쓰는 현대인을 위한 10분 손목 루틴으로 재정의하는 방향이 가장 좋습니다. '
     '또한 기존 상세페이지의 가장 큰 약점은 "누가 왜 사야 하는지"가 첫 화면에서 바로 꽂히지 '
     '않는 점이었습니다. 이번 기획안은 첫 문장부터 [문제 공감 → 제품 차별점 → 안심 보장 → '
     '구매 설득] 순서로 재구성했습니다.', size=9.5, color=GREY, italic=True)
hr()

# =====================================================================
# 1. 제품 최종 포지셔닝
# =====================================================================
heading('1. 제품 최종 포지셔닝', 1)
rich([('제품명  ', {'bold': True, 'color': BLUE}), ('겟머슬 11세대 RPM LED 자이로볼', {})])
rich([('기술명  ', {'bold': True, 'color': BLUE}),
      ('RPM 다이내믹 센서 칩 (RPM Dynamic Sensor Chip)', {})])
rich([('한 줄 정의  ', {'bold': True, 'color': BLUE}),
      ('태엽을 감아 시작하고, 내 손목으로 회전을 유지하며, 회전력에 따라 LED가 달라지는 '
       '손목·악력·전완근 자이로 운동기구', {})])
para('핵심 메시지', bold=True)
shade_para([
    '자동으로 도는 장난감이 아닙니다.',
    '내 손목으로 저항을 키우는 11세대 자이로 운동기구입니다.',
], size=10.5, color=DARK, bold=True)
hr()

# =====================================================================
# 2. 상세페이지 전체 판매 구조
# =====================================================================
heading('2. 상세페이지 전체 판매 구조', 1)
para('단순 기능 나열이 아니라, 고객이 아래 순서대로 설득되도록 설계합니다.', color=GREY)
flow = [
    ('공감', '손목은 매일 쓰는데 따로 운동하지 않는다'),
    ('문제 자각', '손목이 약하면 운동도 일상도 오래 못 버틴다'),
    ('차별점 인식', '구형 랜덤 LED 자이로볼과 다르다'),
    ('기술 호기심', 'RPM 다이내믹 센서 칩으로 회전력 따라 LED가 달라진다'),
    ('운동감 상상', '돌릴수록 저항이 강해지고 손목·악력·전완근이 반응한다'),
    ('구매 명분', '광고비보다 제품력에 집중했다'),
    ('구매 안심', '스트랩·전용케이스 증정 + 30일 무료체험'),
    ('결제 유도', '써보고 별로면 100% 무료 환불이니 지금 사도 손해가 없다'),
]
make_table([['단계', '고객 심리 / 메시지']] + [[f'{i+1}. {k}', v] for i, (k, v) in enumerate(flow)],
           highlight_col=None)
hr()

# =====================================================================
# 3. 상세페이지 15장 구성안
# =====================================================================
heading('3. 상세페이지 15장 구성안', 1)


def scene(no, title_txt, blocks):
    """blocks: list of ('label', value_or_list, kind) ; kind in main/body/emph/img/order"""
    heading(f'{no}장. {title_txt}', 2, color=BLUE)
    for label, val, kind in blocks:
        if kind == 'main':
            para(label, bold=True, color=GREY, size=9, space_after=1)
            shade_para([val] if isinstance(val, str) else val, size=11, color=DARK, bold=True)
        elif kind == 'emph':
            para(label, bold=True, color=GREY, size=9, space_after=1)
            for line in ([val] if isinstance(val, str) else val):
                para(line, bold=True, color=PURPLE)
        elif kind == 'body':
            para(label, bold=True, color=GREY, size=9, space_after=1)
            for line in ([val] if isinstance(val, str) else val):
                para(line)
        elif kind == 'img':
            para(label, bold=True, color=GREY, size=9, space_after=1)
            for line in ([val] if isinstance(val, str) else val):
                bullet(line, color=GREY, size=9.5)
        elif kind == 'order':
            para(label, bold=True, color=GREY, size=9, space_after=1)
            for line in val:
                numbered(line)


scene('1', '메인 후킹', [
    ('목적', '첫 3초 안에 고객을 멈춰 세우기', 'body'),
    ('메인 카피', '손목은 매일 쓰는데, 손목 운동은 왜 안 하세요?', 'main'),
    ('서브 카피', ['마우스, 스마트폰, 헬스, 골프, 테니스까지', '하루 종일 버티는 손목을 위한',
                 '11세대 자이로 손목 운동기구'], 'body'),
    ('강조 문구', ['겟머슬 11세대 RPM LED 자이로볼', '회전력 따라 LED 변화',
                 '손목·악력·전완근 자극', '스트랩 + 전용케이스 증정', '30일 무료체험'], 'emph'),
    ('디자인 방향', ['블랙 배경 + 제품 LED 회전 잔상', '남성 손목 클로즈업 / 전완근 라인 강조',
                  '제품명은 중앙 또는 하단에 고급스럽게 배치'], 'img'),
])

scene('2', '문제 공감', [
    ('메인 카피', '당신의 손목, 생각보다 더 많이 일하고 있습니다', 'main'),
    ('본문 카피', ['헬스할 때 손목이 먼저 흔들린다면', '골프·테니스 후 손목이 쉽게 피로하다면',
                 '마우스와 키보드를 오래 써 손이 뻐근하다면', '일반 악력기가 지루해서 오래 못 했다면',
                 '→ 이제 손목도 따로 운동할 때입니다.'], 'body'),
    ('이미지 방향', ['직장인 / 헬스 / 골프 / 테니스 / 게이머 / 스마트폰 사용자',
                  '6분할 라이프스타일 컷'], 'img'),
])

scene('3', '제품 정의', [
    ('메인 카피', '이건 단순히 반짝이는 공이 아닙니다', 'main'),
    ('본문 카피', ['겟머슬 자이로볼은 손 안에서 회전 저항을 만들고,',
                 '그 저항을 손목으로 버티며 운동하는 자이로 손목 트레이너입니다.'], 'body'),
    ('강한 한 줄', '작다. 간단하다. 그런데 손목은 바로 압니다.', 'emph'),
    ('이미지 방향', ['손바닥 위 제품 컷', '내부 회전 구조 느낌의 3D 그래픽',
                  '"회전 → 저항 → 손목 컨트롤 → 전완근 자극" 흐름도'], 'img'),
])

heading('4장. 구형 제품 저격 비교', 2, color=BLUE)
para('메인 카피', bold=True, color=GREY, size=9, space_after=1)
shade_para(['구형 1~10세대와 비교하지 마세요'], size=11, color=DARK, bold=True)
make_table([
    ['구형 자이로볼의 아쉬움', '겟머슬 11세대의 차이'],
    ['랜덤으로만 번쩍이는 LED', '회전력 따라 LED 변화'],
    ['의미 없는 단순 점멸', 'RPM 다이내믹 센서 칩'],
    ['그냥 돌아가는 장난감 같은 느낌', '손목으로 버티는 자이로 저항감'],
    ['금방 질리는 사용감', 'LED 변화로 운동 몰입감 상승'],
    ['초반 회전 잡기가 번거로움', '태엽식 스타트'],
    ['본품만 있는 구성', '스트랩 + 전용케이스 증정'],
    ['써보기 전 운동감이 불안', '30일 무료체험 보장'],
], highlight_col=1)
rich([('하단 카피  ', {'bold': True, 'color': GREY, 'size': 9}),
      ('구형은 그냥 번쩍였습니다 → 겟머슬은 회전력에 따라 반응합니다', {'bold': True, 'color': PURPLE})])

scene('5', 'RPM 다이내믹 센서 칩', [
    ('메인 카피', '랜덤 LED가 아닙니다', 'main'),
    ('기술명 표기', 'RPM 다이내믹 센서 칩 (RPM Dynamic Sensor Chip)', 'body'),
    ('본문 카피', ['겟머슬 11세대 RPM LED 자이로볼은 회전력에 따라 LED 반응이 달라지는 '
                 'RPM 다이내믹 센서 칩을 적용했습니다.',
                 '느리게 돌릴 때는 은은하게 / 속도가 올라갈수록 선명하게 / 강하게 회전할수록 더 역동적으로'], 'body'),
    ('강한 후킹', '내 손목 스피드가 불빛으로 보입니다', 'emph'),
])
para('⚠️ 주의: 실제 별도 센서칩 구조가 제품에 명확히 들어간 것이 아니라면, 최종 표기는 '
     '"RPM 다이내믹 LED 시스템" 또는 "RPM 반응형 LED 모듈"로 바꾸는 게 안전합니다.',
     size=9, color=RED, italic=True)

heading('6장. 태엽식 스타트', 2, color=BLUE)
para('메인 카피', bold=True, color=GREY, size=9, space_after=1)
shade_para(['태엽 감고 바로 시작하세요'], size=11, color=DARK, bold=True)
para('본문 카피', bold=True, color=GREY, size=9, space_after=1)
para('줄을 찾고, 감고, 실패하는 번거로움 대신 태엽을 감아 초반 회전을 시작합니다. '
     '그다음은 내 손목으로 속도를 키우고 자이로 저항을 버티는 방식입니다.')
para('사용 순서', bold=True, color=GREY, size=9, space_after=1)
for line in ['태엽을 감는다', '회전을 시작한다', '손목 스냅으로 속도를 올린다', 'LED가 살아난다', '저항감이 강해진다']:
    numbered(line)
rich([('하단 카피  ', {'bold': True, 'color': GREY, 'size': 9}),
      ('시작은 쉽게, 운동감은 직접적으로', {'bold': True, 'color': PURPLE})])

scene('7', '진짜 운동감 강조', [
    ('메인 카피', '자동으로 도는 장난감이 아닙니다', 'main'),
    ('본문 카피', ['겟머슬은 스스로 계속 돌아가는 제품이 아닙니다.',
                 '내 손목 움직임으로 회전을 유지하고 내 힘으로 저항을 키우는 진짜 자이로 운동기구입니다.'], 'body'),
    ('강조 카피', ['내가 돌린 만큼 저항이 커지고, 내가 컨트롤한 만큼 손목이 반응합니다.'], 'emph'),
    ('감각 카피', ['회전 소리까지 느껴지는 운동감', '손 안에서 차오르는 묵직한 자이로 저항',
                 '돌릴수록 살아나는 손목·전완근 자극'], 'img'),
])

heading('8장. 악력기·덤벨과 비교', 2, color=BLUE)
para('메인 카피', bold=True, color=GREY, size=9, space_after=1)
shade_para(['악력기는 쥐기만 합니다 — 자이로볼은 손목이 계속 반응합니다'], size=11, color=DARK, bold=True)
make_table([
    ['구분', '악력기', '덤벨', '겟머슬 자이로볼'],
    ['운동 방식', '쥐기 반복', '무게 들기', '회전 저항 유지'],
    ['재미', '단조로움', '보통', 'LED + 회전감'],
    ['휴대성', '좋음', '낮음', '좋음'],
    ['손목 컨트롤', '낮음', '보통', '높음'],
    ['전완근 자극', '가능', '가능', '회전 유지형 자극'],
    ['몰입감', '낮음', '보통', '높음'],
], highlight_col=3)
rich([('하단 카피  ', {'bold': True, 'color': GREY, 'size': 9}),
      ('쥐는 힘만이 아니라 손목으로 버티는 힘까지', {'bold': True, 'color': PURPLE})])

scene('9', '손목·악력·전완근 자극', [
    ('메인 카피', '손목부터 전완근까지, 손 안에서 자극', 'main'),
    ('본문 카피', ['회전을 유지할수록 자이로 저항감이 강해집니다.',
                 '그 저항을 손목으로 버티며 손목·악력·전완근 운동에 도움을 줄 수 있습니다.'], 'body'),
    ('강조 부위', ['손목 / 악력 / 전완근 / 팔 회전 감각'], 'emph'),
    ('디자인 방향', ['손목에서 전완근으로 이어지는 네온 라인', '근육 자극 부위 도해',
                  '의료 느낌 말고 스포츠 장비 느낌'], 'img'),
])

heading('10장. 타깃별 사용 장면', 2, color=BLUE)
para('메인 카피', bold=True, color=GREY, size=9, space_after=1)
shade_para(['손목을 많이 쓰는 사람이라면 하나쯤 필요합니다'], size=11, color=DARK, bold=True)
make_table([
    ['타깃', '제안 루틴'],
    ['헬스', '덤벨·바벨·턱걸이 전 손목과 전완근을 깨우는 루틴'],
    ['골프', '스윙 전 손목 감각과 전완근 컨트롤 루틴'],
    ['테니스·배드민턴', '라켓을 잡는 손목 힘과 회전 감각 관리'],
    ['클라이밍', '손가락·손목·전완근을 함께 쓰는 감각 훈련'],
    ['직장인·게이머', '마우스와 키보드로 지친 손목을 위한 짧은 루틴'],
    ['선물용', '운동 좋아하는 사람에게 부담 없는 실용 선물'],
])

heading('11장. 하루 10분 루틴', 2, color=BLUE)
para('메인 카피', bold=True, color=GREY, size=9, space_after=1)
shade_para(['처음이라면 하루 1분부터 시작하세요'], size=11, color=DARK, bold=True)
make_table([
    ['단계', '시간', '방법', '목적'],
    ['Lv.1 적응', '30초~1분', '천천히 회전 유지', '감각 익히기'],
    ['Lv.2 유지', '2~3분', '일정 속도 유지', '손목 컨트롤'],
    ['Lv.3 강화', '3~5분', '속도 올리기', '전완근 자극'],
    ['Lv.4 루틴', '5~10분', '양손 번갈아 사용', '홈트 루틴'],
])
para('처음부터 오래 하지 마세요. 짧게 시작하고, 익숙해질수록 속도와 시간을 늘려보세요.', color=GREY)

scene('12', '단일 구성 + 사은품 강조', [
    ('메인 카피', '옵션 고민 없이, 필요한 구성만 한 번에', 'main'),
    ('구성품', ['겟머슬 11세대 RPM LED 자이로볼 본품', '안전 스트랩 증정', '전용 보관 케이스 증정'], 'emph'),
    ('본문 카피', ['고속 회전 시 안정감을 더해주는 스트랩, 보관과 휴대를 편하게 해주는 전용 케이스까지.',
                 '사자마자 바로 쓰고, 쓰고 나면 깔끔하게 보관하세요.'], 'body'),
    ('배너 카피', ['본품 + 스트랩 + 전용케이스', '단일 구성으로 바로 시작'], 'img'),
])
para('⚠️ 주의: 색상 옵션은 전부 제외 · 컬러 선택 문구 금지 · 단일 제품으로만 구성',
     size=9, color=RED, italic=True)

scene('13', '광고비 대신 제품력', [
    ('메인 카피', '비싼 자이로볼 가격, 정말 제품값일까요 광고값일까요?', 'main'),
    ('본문 카피', ['겟머슬은 유명인 광고비를 제품 가격에 얹지 않았습니다.',
                 '그 비용 대신 11세대 구조 / 태엽식 스타트 / RPM 다이내믹 센서 칩 / '
                 '자이로 회전 저항감 / 스트랩과 전용 케이스 구성에 집중했습니다.'], 'body'),
    ('강조 카피', '광고비는 빼고, 제품력은 채운 11세대 자이로볼', 'emph'),
])

scene('14', '30일 무료체험 보장', [
    ('메인 카피', '뜯어서 30일 동안 마음껏 돌려보세요', 'main'),
    ('서브 카피', '제품에 자신 있으니까 약속합니다', 'body'),
    ('체크 포인트', ['개봉해도 OK', '사용해도 OK', '단순변심도 OK'], 'emph'),
    ('최종 후킹', '만족하지 못하면 100% 무료 환불', 'emph'),
    ('본문 카피', ['사진만 보고는 자이로볼의 운동감을 알 수 없습니다.',
                 '직접 돌려보고, LED 반응을 보고, 손목에 전해지는 저항감을 느껴보세요.',
                 '만족하지 못하셨다면 30일 이내 100% 무료 환불해드립니다.'], 'body'),
])
para('※ 30일 무료체험 및 무료 환불은 상품 수령일 기준 30일 이내 신청 건에 한해 적용됩니다. '
     '구성품 누락, 고의 훼손, 사용 불가 수준의 파손이 있는 경우 제한될 수 있습니다.',
     size=8.5, color=GREY)

heading('15장. FAQ + 최종 구매 클로징', 2, color=BLUE)
para('메인 카피', bold=True, color=GREY, size=9, space_after=1)
shade_para(['구매 전 궁금한 점, 미리 확인하세요'], size=11, color=DARK, bold=True)
faqs = [
    ('Q. 소리가 나나요?',
     '네. 자이로볼은 내부 회전체가 고속으로 회전하는 운동기구 특성상 회전 소리와 진동이 발생할 수 '
     '있습니다. 이는 제품의 정상적인 작동 특성입니다.'),
    ('Q. 자동으로 계속 돌아가나요?',
     '아니요. 태엽으로 초반 회전을 시작한 뒤, 사용자의 손목 움직임으로 회전을 유지하는 제품입니다.'),
    ('Q. 초보자도 사용할 수 있나요?',
     '처음에는 30초~1분씩 천천히 시작하는 것을 권장합니다. 익숙해질수록 회전 유지 감각이 좋아지고 '
     '더 강한 저항감을 느낄 수 있습니다.'),
    ('Q. LED는 어떻게 나오나요?',
     '회전력과 속도에 따라 LED 반응이 달라지는 구조입니다. 손목 스피드에 따라 LED 밝기와 연출감이 '
     '다르게 나타날 수 있습니다.'),
    ('Q. 손목 통증 치료용인가요?',
     '아니요. 본 제품은 의료기기가 아닌 운동 보조용품입니다. 통증이나 질환이 있는 경우 전문가 상담 '
     '후 사용해 주세요.'),
]
for q, a in faqs:
    para(q, bold=True, color=BLUE, space_after=1)
    para(a, space_after=6)
para('최종 클로징', bold=True, color=GREY, size=9, space_after=1)
shade_para([
    '저가형 자이로볼에 실망했다면, 이제 11세대로 바꾸세요',
    '',
    '랜덤 LED가 아닙니다. 회전력에 따라 LED가 달라집니다.',
    '자동 장난감이 아닙니다. 내 손목으로 저항을 키우는 운동기구입니다.',
    '광고비로 가격을 부풀리지 않았습니다. 제품력, 사은품, 30일 무료체험으로 증명합니다.',
], fill='111418', border='111418', size=10.5, color=RGBColor(0xFF, 0xFF, 0xFF), bold=True)
rich([('최종 CTA  ', {'bold': True, 'color': GOLD}),
      ('광고비는 빼고, 제품력은 채우고, 30일 무료체험까지 건 11세대 자이로볼 — '
       '겟머슬로 손목 운동 루틴을 시작하세요', {'bold': True})])
hr()

# =====================================================================
# 4. 썸네일 문구 최종안
# =====================================================================
heading('4. 썸네일 문구 최종안', 1)
thumbs = [
    ('1안 — 가장 추천', ['회전력 따라 LED 변화', '11세대 RPM LED 자이로볼', '손목·악력·전완근 자극'], True),
    ('2안 — 구형 저격형', ['랜덤 LED 구형은 그만', 'RPM 다이내믹 센서 칩', '겟머슬 11세대'], False),
    ('3안 — 보장 강조형', ['30일 무료체험', '개봉·사용·단순변심 OK', '100% 무료 환불'], False),
    ('4안 — 구성품 강조형', ['스트랩 + 전용케이스 증정', '단일 구성 바로 사용', '겟머슬 자이로볼'], False),
    ('5안 — 가격 명분형', ['광고비 뺀 가격', '제품력 채운 11세대', '30일 무료체험'], False),
]
for name, lines, rec in thumbs:
    rich([(name, {'bold': True, 'color': (GOLD if rec else BLUE)})], space_after=1)
    shade_para(lines, fill=('FFF6E0' if rec else 'F4F4F6'),
               border=('E5C76B' if rec else 'DDDDDD'), size=10, color=DARK, bold=True)
hr()

# =====================================================================
# 5. 첫 화면 최종 조합
# =====================================================================
heading('5. 상세페이지 첫 화면 최종 조합', 1)
para('디자이너에게 첫 화면은 아래 그대로 넣게 하세요.', color=GREY)
shade_para([
    '손목은 매일 쓰는데, 손목 운동은 왜 안 하세요?',
    '',
    '겟머슬 11세대 RPM LED 자이로볼',
    '회전력 따라 LED가 달라지는 RPM 다이내믹 센서 칩 적용',
    '돌릴수록 강해지는 자이로 저항감으로 손목·악력·전완근 자극',
], size=10.5, color=DARK, bold=True)
para('하단 배지', bold=True, color=GREY, size=9, space_after=1)
for t in ['스트랩 증정', '전용케이스 증정', '30일 무료체험', '만족 못하면 100% 무료 환불']:
    bullet(t, color=GREEN, bold=True)
hr()

# =====================================================================
# 6. 반복 핵심 문장
# =====================================================================
heading('6. 페이지 전반에서 반복할 핵심 문장', 1)
para('아래 문장은 페이지 중간중간 반복해서 박아야 합니다.', color=GREY)
for t in [
    '랜덤 LED 구형은 그만',
    '회전력 따라 LED가 달라지는 11세대 자이로볼',
    '자동으로 도는 장난감이 아닙니다',
    '내 손목으로 저항을 키우는 운동기구입니다',
    '속도감은 눈으로, 저항감은 손목으로',
    '악력기는 쥐기만 합니다. 자이로볼은 손목이 계속 반응합니다',
    '광고비는 빼고, 제품력은 채웠습니다',
    '뜯어서 30일 동안 마음껏 돌려보세요',
    '개봉해도 OK, 사용해도 OK, 단순변심도 OK',
    '만족하지 못하면 100% 무료 환불',
]:
    bullet(t, bold=True)
hr()

# =====================================================================
# 7 / 8. 금지 표현 & 대체 표현
# =====================================================================
heading('7. 절대 쓰면 안 되는 표현', 1)
for t in [
    'AI 칩 / 인공지능 LED / AI가 회전 분석',
    '무소음 / 소음 없음 / 완전 무진동',
    '자동 회전 보정 / 자동 궤도 보정',
    '고장 절대 없음 / 초보자도 무조건 100% 가능',
    '손목 통증 치료 / 재활 치료 / 관절 개선 / 염증 완화',
    '근육 증가 보장 / 악력 무조건 증가',
    '의료기기처럼 보이는 표현',
]:
    bullet(t, color=RED)

heading('8. 대신 써야 하는 표현', 1)
for t in [
    'RPM 다이내믹 센서 칩 / RPM Dynamic Sensor Chip',
    'RPM 반응형 LED / 회전력 반응 LED / 회전 단계별 LED / 다이내믹 LED 퍼포먼스',
    '돌릴수록 강해지는 자이로 저항감',
    '내 손목으로 직접 유지하는 회전 운동',
    '손목·악력·전완근 자극',
    '자이로볼 특유의 회전 소리와 진동',
    '손 안에서 시작하는 10분 손목 루틴',
    '운동 보조용품 / 손목 컨트롤 훈련 / 전완근 루틴',
]:
    bullet(t, color=GREEN)
hr()

# =====================================================================
# 9. 쿠팡 상품명 추천
# =====================================================================
heading('9. 쿠팡 상품명 추천', 1)
rich([('검색 최적화형  ', {'bold': True, 'color': BLUE})], space_after=1)
para('겟머슬 11세대 RPM LED 자이로볼 태엽식 손목운동기구 전완근 악력 운동 파워볼 스트랩 케이스 세트')
rich([('전환형  ', {'bold': True, 'color': BLUE})], space_after=1)
para('겟머슬 11세대 RPM LED 자이로볼 손목 악력 전완근 운동기구 스트랩 케이스 포함')
rich([('브랜드형  ', {'bold': True, 'color': BLUE})], space_after=1)
para('겟머슬 11세대 RPM LED 자이로볼')
hr()

# =====================================================================
# 10. 최종 구매전환 카피
# =====================================================================
heading('10. 최종 구매전환 카피', 1)
shade_para([
    'AI라고 과장하지 않습니다. 광고비로 가격을 부풀리지도 않았습니다.',
    '',
    '겟머슬 11세대 RPM LED 자이로볼은 태엽식 스타트, RPM 다이내믹 센서 칩,',
    '자이로 저항감에 집중한 손목·악력·전완근 운동기구입니다.',
    '',
    '여기에 스트랩과 전용케이스를 기본 증정하고,',
    '제품력에 자신 있으니까 30일 무료체험까지 약속합니다.',
    '',
    '개봉해도 OK · 사용해도 OK · 단순변심도 OK',
    '직접 돌려보고 만족하지 못하셨다면 100% 무료 환불해드립니다.',
    '',
    '겟머슬의 자신감이 진짜인지, 손목으로 직접 확인하세요.',
], fill='111418', border='111418', size=10.5, color=RGBColor(0xFF, 0xFF, 0xFF), bold=True)

# =====================================================================
out = os.path.join(os.path.dirname(__file__), '..', 'docs', '겟머슬_상세페이지_기획안.docx')
out = os.path.abspath(out)
doc.save(out)
print('saved:', out)
